"""Migração da planilha para o banco: casa cada linha com uma conta e registra o que falta.

Ordem de casamento: raiz do CNPJ, depois domínio, depois nome normalizado (só quando não há
CNPJs diferentes em jogo). Campos preenchidos na planilha atualizam a conta; vazios não apagam.
CNPJ e domínio só entram onde ainda não havia; divergência vira buraco para revisão.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path

import psycopg

from .identidade import (
    cnpj_normalizado,
    cnpj_raiz,
    dominio_generico,
    e_matriz,
    normalizar_chave,
    normalizar_dominio,
    normalizar_nome,
    so_digitos,
)

ALIASES = {
    "nome": ["empresa", "nome", "conta", "razao_social", "nome_fantasia", "company", "account", "organizacao"],
    "cnpj": ["cnpj", "cnpj_matriz"],
    "site": ["site", "dominio", "website", "url", "domain", "site_empresa"],
    "uf": ["uf", "estado", "state"],
    "cidade": ["cidade", "municipio", "city"],
    "setor": ["setor", "segmento", "industria", "industry", "vertical"],
    "braco": ["braco", "braco_icp", "icp", "arm"],
    "abc": ["abc", "classe", "curva_abc", "tier", "classificacao"],
    "grupo": ["grupo", "grupo_economico", "holding"],
    "decisor": ["decisor", "contato", "decisor_sugerido", "nome_decisor"],
    "cargo": ["cargo", "cargo_decisor", "titulo", "title"],
    "email": ["email", "email_decisor", "e_mail"],
    "linkedin": ["linkedin", "linkedin_decisor", "perfil_linkedin"],
    "uf_decisor": ["uf_decisor", "estado_decisor"],
    "headcount": ["headcount", "funcionarios", "colaboradores", "employees", "numero_de_funcionarios"],
    "capital_social": ["capital_social", "capital"],
    "observacoes": ["observacoes", "obs", "notas", "notes"],
}


def ler_planilha(caminho: str | Path) -> list[dict[str, str]]:
    """Lê CSV (vírgula, ponto e vírgula ou tabulação) ou XLSX (primeira aba) com cabeçalho na 1ª linha."""
    caminho = Path(caminho)
    if caminho.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        aba = load_workbook(caminho, read_only=True, data_only=True).worksheets[0]
        linhas = [["" if v is None else str(v).strip() for v in linha] for linha in aba.iter_rows(values_only=True)]
    else:
        texto = caminho.read_bytes().decode("utf-8-sig", errors="replace")
        primeira = texto.split("\n", 1)[0]
        delim = max(["\t", ";", ","], key=primeira.count)
        linhas = [[c.strip() for c in linha] for linha in csv.reader(io.StringIO(texto), delimiter=delim)]
    linhas = [l for l in linhas if any(l)]
    if not linhas:
        return []
    cabecalho = [normalizar_chave(c) for c in linhas[0]]
    return [dict(zip(cabecalho, l + [""] * (len(cabecalho) - len(l)))) for l in linhas[1:]]


def pegar(linha: dict[str, str], campo: str) -> str:
    for alias in ALIASES[campo]:
        valor = (linha.get(alias) or "").strip()
        if valor:
            return valor
    return ""


@dataclass
class Relatorio:
    linhas: int = 0
    novas: int = 0
    atualizadas: int = 0
    pessoas: int = 0
    estabelecimentos: int = 0
    avisos: list[str] = field(default_factory=list)


def registrar_buraco(conn, conta_id, tipo: str, detalhe: str = "", linha: int | None = None) -> None:
    conn.execute(
        """insert into buraco (conta_id, tipo, detalhe, linha_origem) values (%s, %s, %s, %s)
           on conflict (conta_id, tipo) where aberto and conta_id is not null do update set detalhe = excluded.detalhe""",
        (conta_id, tipo, detalhe, linha),
    )


def _achar_conta(conn, raiz, dominio, nome_norm):
    if raiz:
        c = conn.execute("select * from conta where cnpj_raiz = %s", (raiz,)).fetchone()
        if c:
            return c, "cnpj"
    if dominio:
        c = conn.execute("select * from conta where dominio = %s", (dominio,)).fetchone()
        if c and not (raiz and c["cnpj_raiz"] and c["cnpj_raiz"] != raiz):
            return c, "dominio"
    if nome_norm:
        candidatas = conn.execute("select * from conta where nome_normalizado = %s", (nome_norm,)).fetchall()
        compativeis = [c for c in candidatas if not (raiz and c["cnpj_raiz"] and c["cnpj_raiz"] != raiz)]
        if len(compativeis) == 1:
            return compativeis[0], "nome"
    return None, None


def importar(conn: psycopg.Connection, linhas: list[dict[str, str]], origem: str = "planilha") -> Relatorio:
    rel = Relatorio()
    with conn.transaction():
        for i, linha in enumerate(linhas):
            numero = i + 2
            rel.linhas += 1
            nome = pegar(linha, "nome")
            cnpj_bruto = pegar(linha, "cnpj")
            if not nome and not cnpj_bruto:
                registrar_buraco(conn, None, "linha_vazia", "linha sem nome nem CNPJ", numero)
                continue
            cnpj = cnpj_normalizado(cnpj_bruto)
            raiz = cnpj_raiz(cnpj)
            email = pegar(linha, "email")
            dominio = normalizar_dominio(pegar(linha, "site")) or normalizar_dominio(email if "@" in email else "")
            generico = dominio_generico(dominio)
            if generico:
                dominio = None
            nome_norm = normalizar_nome(nome) if nome else ""

            conta, _ = _achar_conta(conn, raiz, dominio, nome_norm)
            dados = {
                "uf": pegar(linha, "uf").upper()[:2] or None,
                "cidade": pegar(linha, "cidade") or None,
                "setor": pegar(linha, "setor") or None,
                "braco": pegar(linha, "braco") or None,
                "abc": (pegar(linha, "abc").upper()[:1] or None),
                "headcount": int(so_digitos(pegar(linha, "headcount"))) if so_digitos(pegar(linha, "headcount")) else None,
                "observacoes": pegar(linha, "observacoes") or None,
            }
            if dados["abc"] not in (None, "A", "B", "C"):
                rel.avisos.append(f"Linha {numero}: classe ABC desconhecida ({pegar(linha, 'abc')})")
                dados["abc"] = None

            if conta is None:
                dono_dominio = dominio and conn.execute("select id from conta where dominio = %s", (dominio,)).fetchone()
                conta = conn.execute(
                    """insert into conta (nome, nome_normalizado, cnpj_raiz, dominio, origem, uf, cidade, setor, braco, abc, headcount, observacoes)
                       values (%(nome)s, %(nome_norm)s, %(raiz)s, %(dominio)s, %(origem)s, %(uf)s, %(cidade)s, %(setor)s, %(braco)s, %(abc)s, %(headcount)s, %(observacoes)s)
                       returning *""",
                    {
                        **dados,
                        "nome": nome or f"CNPJ {cnpj_bruto}",
                        "nome_norm": nome_norm or so_digitos(cnpj_bruto),
                        "raiz": raiz,
                        "dominio": None if dono_dominio else dominio,
                        "origem": origem,
                    },
                ).fetchone()
                rel.novas += 1
                if dono_dominio:
                    registrar_buraco(
                        conn, conta["id"], "conflito_cnpj_dominio",
                        f"o domínio {dominio} já pertence a outra conta com CNPJ diferente", numero,
                    )
            else:
                sets = {k: v for k, v in dados.items() if v is not None}
                if nome:
                    sets["nome"] = nome
                    sets["nome_normalizado"] = nome_norm
                if raiz and not conta["cnpj_raiz"]:
                    sets["cnpj_raiz"] = raiz
                if dominio and not conta["dominio"]:
                    dono = conn.execute("select id from conta where dominio = %s", (dominio,)).fetchone()
                    if dono:
                        registrar_buraco(conn, conta["id"], "conflito_cnpj_dominio", f"o domínio {dominio} já pertence a outra conta", numero)
                    else:
                        sets["dominio"] = dominio
                elif dominio and conta["dominio"] != dominio:
                    registrar_buraco(
                        conn, conta["id"], "conflito_cnpj_dominio",
                        f"a planilha traz {dominio}, a conta já tem {conta['dominio']}", numero,
                    )
                if sets:
                    colunas = ", ".join(f"{k} = %({k})s" for k in sets)
                    conn.execute(f"update conta set {colunas}, atualizado_em = now() where id = %(id)s", {**sets, "id": conta["id"]})
                rel.atualizadas += 1

            if cnpj_bruto and not cnpj:
                registrar_buraco(conn, conta["id"], "cnpj_invalido", f"CNPJ da planilha: {cnpj_bruto}", numero)
            if generico:
                registrar_buraco(conn, conta["id"], "dominio_generico", "a planilha só traz e-mail ou site genérico", numero)
            if cnpj:
                feito = conn.execute(
                    """insert into estabelecimento (cnpj, conta_id, matriz, uf, cidade) values (%s, %s, %s, %s, %s)
                       on conflict (cnpj) do nothing returning cnpj""",
                    (cnpj, conta["id"], e_matriz(cnpj), dados["uf"], dados["cidade"]),
                ).fetchone()
                rel.estabelecimentos += bool(feito)

            grupo = pegar(linha, "grupo")
            if grupo:
                g = conn.execute("select id from grupo_economico where nome = %s", (grupo,)).fetchone() or conn.execute(
                    "insert into grupo_economico (nome) values (%s) returning id", (grupo,)
                ).fetchone()
                conn.execute("update conta set grupo_economico_id = %s where id = %s", (g["id"], conta["id"]))

            decisor = pegar(linha, "decisor")
            if decisor:
                conn.execute(
                    """insert into pessoa (conta_id, nome, cargo, papel_comite, email, linkedin, uf, origem)
                       values (%s, %s, %s, 'decisor', %s, %s, %s, %s)
                       on conflict (conta_id, nome) do update set
                         cargo = coalesce(excluded.cargo, pessoa.cargo),
                         email = coalesce(excluded.email, pessoa.email),
                         linkedin = coalesce(excluded.linkedin, pessoa.linkedin),
                         uf = coalesce(excluded.uf, pessoa.uf)""",
                    (
                        conta["id"], decisor, pegar(linha, "cargo") or None, email or None,
                        pegar(linha, "linkedin") or None, pegar(linha, "uf_decisor").upper()[:2] or None, origem,
                    ),
                )
                rel.pessoas += 1
        atualizar_buracos(conn)
    return rel


def atualizar_buracos(conn: psycopg.Connection) -> None:
    """Abre os buracos de identidade que faltam e fecha os que já foram preenchidos."""
    conn.execute(
        """insert into buraco (conta_id, tipo, detalhe)
           select id, 'sem_cnpj', 'conta sem CNPJ válido' from conta where cnpj_raiz is null
           on conflict (conta_id, tipo) where aberto and conta_id is not null do nothing"""
    )
    conn.execute(
        """insert into buraco (conta_id, tipo, detalhe)
           select id, 'sem_dominio', 'conta sem domínio' from conta where dominio is null
           on conflict (conta_id, tipo) where aberto and conta_id is not null do nothing"""
    )
    conn.execute(
        """update buraco b set aberto = false, resolvido_em = now(), resolucao = coalesce(b.resolucao, 'preenchido')
           from conta c where b.conta_id = c.id and b.aberto and (
             (b.tipo in ('sem_cnpj', 'cnpj_invalido') and c.cnpj_raiz is not null) or
             (b.tipo in ('sem_dominio', 'dominio_generico') and c.dominio is not null))"""
    )
    # Mesmo nome com CNPJs diferentes e fora de um grupo comum: filial cadastrada à parte ou grupo econômico.
    conn.execute(
        """insert into buraco (conta_id, tipo, detalhe)
           select c.id, 'possivel_grupo', 'mesmo nome de ' || string_agg(o.nome || coalesce(' (' || o.cnpj_raiz || ')', ''), ', ')
           from conta c join conta o on o.nome_normalizado = c.nome_normalizado and o.id <> c.id
           where c.grupo_economico_id is null or c.grupo_economico_id is distinct from o.grupo_economico_id
           group by c.id
           on conflict (conta_id, tipo) where aberto and conta_id is not null do nothing"""
    )
    conn.execute(
        """update buraco b set aberto = false, resolvido_em = now(), resolucao = coalesce(b.resolucao, 'resolvido')
           where b.aberto and b.tipo = 'possivel_grupo' and not exists (
             select 1 from conta c join conta o on o.nome_normalizado = c.nome_normalizado and o.id <> c.id
             where c.id = b.conta_id and (c.grupo_economico_id is null or c.grupo_economico_id is distinct from o.grupo_economico_id))"""
    )
